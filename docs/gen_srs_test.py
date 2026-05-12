
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT      = "Times New Roman"
SIZE_BODY = Pt(12)
SIZE_H1   = Pt(14)
SIZE_H2   = Pt(13)

doc = Document()
sec = doc.sections[0]
sec.left_margin   = Cm(3.0)
sec.right_margin  = Cm(2.0)
sec.top_margin    = Cm(2.5)
sec.bottom_margin = Cm(2.5)

sn = doc.styles["Normal"]
sn.font.name = FONT
sn.font.size = SIZE_BODY
sn.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
sn.paragraph_format.space_after  = Pt(0)
sn.paragraph_format.space_before = Pt(0)

for hn, sz in [("Heading 1", SIZE_H1), ("Heading 2", SIZE_H2), ("Heading 3", Pt(12))]:
    s = doc.styles[hn]
    s.font.name  = FONT; s.font.size = sz; s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(18); s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

footer = sec.footer
fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; fp.clear()
run = fp.add_run(); run.font.name = FONT; run.font.size = Pt(10)
for tag in [("begin", None), (None, "PAGE"), ("end", None)]:
    if tag[0]: fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag[0]); run._r.append(fc)
    else: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = tag[1]; run._r.append(it)

def h(text, level=1):
    p = doc.add_heading(text, level=level); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs: r.font.name = FONT; r.font.color.rgb = RGBColor(0,0,0)

def para(text, fi=True):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(8); p.paragraph_format.space_before = Pt(0)
    if fi: p.paragraph_format.first_line_indent = Cm(1.25)
    r = p.add_run(text); r.font.name = FONT; r.font.size = SIZE_BODY; return p

def bul(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text); r.font.name = FONT; r.font.size = SIZE_BODY; return p

def lbl(label, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Cm(1.25); p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(label + ": "); r1.bold = True; r1.font.name = FONT; r1.font.size = SIZE_BODY
    r2 = p.add_run(text); r2.font.name = FONT; r2.font.size = SIZE_BODY; return p

def rq(rid, prio, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Cm(1.25); p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(f"[{rid}] ({prio})  "); r1.bold = True; r1.font.name = FONT; r1.font.size = SIZE_BODY
    r2 = p.add_run(text); r2.font.name = FONT; r2.font.size = SIZE_BODY; return p

def cp(text, sz=12, bold=False, sb=0, sa=6):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run(text); r.bold = bold; r.font.name = FONT; r.font.size = Pt(sz); return p

print("helpers ok")
